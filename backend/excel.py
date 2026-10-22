from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re

from openpyxl import load_workbook

HEADERS = (
    "SKU编号", "SKU名称", "账号", "时间", "货架位", "原库存", "变动",
    "新库存", "仓库", "原成本价", "出入库价", "现成本价", "类型", "单据号",
)
MAX_RECORDS = 50000
MAX_BYTES = 20 * 1024 * 1024


def number(value):
    text = re.sub(r"[\s,，]", "", str(value)).replace("−", "-")
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", text):
        return None
    try:
        result = Decimal(text)
        return result if result.is_finite() and abs(result) < Decimal("1e26") else None
    except InvalidOperation:
        return None


def cell_text(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        fmt = cell.number_format or "General"
        if re.fullmatch(r"0+", fmt):
            return str(int(value)).zfill(len(fmt)) if value == int(value) else str(value)
        match = re.fullmatch(r"(?:#,##)?0\.(0+)", fmt)
        if match:
            return f"{value:.{len(match[1])}f}"
        return str(value) if isinstance(value, int) else format(value, ".15g")
    return str(value)


def parse_excel(content: bytes):
    if not content or len(content) > MAX_BYTES:
        raise ValueError("请选择不超过 20 MB 的 Excel 文件")
    if not content.startswith(b"PK"):
        raise ValueError("文件内容不是有效的 .xlsx 工作簿")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except Exception:
        raise ValueError("无法读取 Excel 文件，请确认文件完整且未设置密码") from None
    sheets = []
    skipped = []
    total = 0
    try:
        for worksheet in workbook.worksheets:
            if worksheet.max_column > 200 or worksheet.max_row > MAX_RECORDS + 21:
                raise ValueError("工作表超出支持范围：最多 200 列、50,000 条记录")
            indices = None
            records = []
            best_headers = []
            best_score = -1
            for row_number, row in enumerate(worksheet.iter_rows(), start=1):
                if indices is None:
                    if row_number > 20:
                        break
                    headers = [re.sub(r"[\s\ufeff]", "", cell_text(cell)) for cell in row]
                    score = sum(header in headers for header in HEADERS)
                    if score > best_score:
                        best_score, best_headers = score, headers
                    if score != len(HEADERS):
                        continue
                    if any(headers.count(header) != 1 for header in HEADERS):
                        raise ValueError(f"工作表 {worksheet.title} 存在重复表头")
                    indices = [headers.index(header) for header in HEADERS]
                    continue
                cells = [cell_text(row[index]) if index < len(row) else "" for index in indices]
                if not any(value.strip() for value in cells):
                    continue
                if any(len(value) > 10000 for value in cells):
                    raise ValueError(f"第 {row_number} 行单元格内容过长")
                for index in (0, 2, 3, 4, 8, 12, 13):
                    if len(cells[index]) > 255:
                        raise ValueError(f"第 {row_number} 行的 {HEADERS[index]} 超过 255 个字符")
                records.append({"rowNumber": row_number, "cells": cells})
                total += 1
                if total > MAX_RECORDS:
                    raise ValueError("每个文件最多支持 50,000 条记录，请拆分后上传")
            if indices is None:
                missing = [header for header in HEADERS if header not in best_headers]
                skipped.append({"name": worksheet.title, "reason": "缺少表头：" + "、".join(missing)})
            elif not records:
                skipped.append({"name": worksheet.title, "reason": "只有表头，没有数据记录"})
            else:
                sheets.append({"name": worksheet.title, "records": records})
    finally:
        workbook.close()
    if not sheets:
        reason = skipped[0]["reason"] if len(skipped) == 1 else "未找到符合库存更新记录格式的工作表"
        raise ValueError(reason)
    return {"sheets": sheets, "skippedSheets": skipped, "recordCount": total}
