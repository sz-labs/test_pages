"""Local functional coverage. All writes use disposable SQLite databases."""

from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import importlib
from pathlib import Path
import time
from zipfile import ZIP_STORED, ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, event, func, select

from backend.app import create_app
from backend.excel import HEADERS, MAX_BYTES, MAX_RECORDS, parse_excel
from backend.repository import InventoryRepository, records, uploads
from backend.settings import Settings
from backend.tests.test_api import PASSWORD, api, auth, sample_row, workbook_bytes


def padded_workbook(size):
    content = workbook_bytes()
    name = "boundary-padding.bin"
    padding = size - len(content) - 30 - 46 - 2 * len(name)
    output = BytesIO(content)
    with ZipFile(output, "a", compression=ZIP_STORED) as archive:
        archive.writestr(name, b"\0" * padding)
    assert len(output.getvalue()) == size
    return output.getvalue()


def test_health_missing_records_and_invalid_uuid(api):
    client, _ = api
    response = client.get("/api/health")
    assert response.json() == {"status": "ok", "database": "connected"}
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert client.get("/api/uploads/00000000-0000-4000-8000-000000000001").status_code == 404
    assert client.get("/api/uploads/not-a-uuid").status_code == 422


@pytest.mark.parametrize("filename,content,status", [
    ("inventory.csv", b"not excel", 400),
    ("inventory.xlsx", b"", 422),
    ("inventory.xlsx", b"PKbroken", 422),
    ("x" * 251 + ".xlsx", b"invalid", 400),
])
def test_upload_validation_does_not_publish(api, filename, content, status):
    client, _ = api
    response = client.post("/api/uploads", headers=auth(client), files={"file": (filename, content)})
    assert response.status_code == status
    assert client.get("/api/uploads").json()["items"] == []


def test_exact_20_mib_upload_and_one_byte_over_limit(api):
    client, _ = api
    headers = auth(client)
    content = padded_workbook(MAX_BYTES)
    saved = client.post("/api/uploads", headers=headers, files={"file": ("boundary.xlsx", content)})
    assert saved.status_code == 201
    assert saved.json()["recordCount"] == 1
    oversized = client.post("/api/uploads", headers=headers, files={"file": ("too-large.xlsx", content + b"0")})
    assert oversized.status_code == 413
    assert len(client.get("/api/uploads").json()["items"]) == 1


def test_request_header_size_and_empty_upload_are_rejected(api):
    client, _ = api
    headers = auth(client)
    assert client.post("/api/uploads", headers={**headers, "content-length": str(MAX_BYTES + 1024 * 1024 + 1)}).status_code == 413
    assert client.post("/api/uploads", headers={**headers, "content-length": "invalid"}).status_code == 400
    assert client.post("/api/uploads", headers=headers).status_code == 422


def test_eight_hour_expired_token_is_rejected(api, monkeypatch):
    client, _ = api
    response = client.post("/api/session", json={"password": PASSWORD})
    assert response.json()["expiresIn"] == 8 * 3600
    headers = {"Authorization": "Bearer " + response.json()["accessToken"]}
    app_module = importlib.import_module("backend.app")
    now = time.time()
    monkeypatch.setattr(app_module.time, "time", lambda: now + 8 * 3600 + 2)
    assert client.post("/api/uploads", headers=headers, files={"file": ("sample.xlsx", workbook_bytes())}).status_code == 401


def test_ten_failed_logins_rate_limit_then_recover(api, monkeypatch):
    client, _ = api
    for _ in range(10):
        assert client.post("/api/session", json={"password": "incorrect"}).status_code == 401
    assert client.post("/api/session", json={"password": PASSWORD}).status_code == 429
    app_module = importlib.import_module("backend.app")
    later = time.monotonic() + 301
    monkeypatch.setattr(app_module.time, "monotonic", lambda: later)
    assert client.post("/api/session", json={"password": PASSWORD}).status_code == 200


def test_disallowed_origin_cannot_read_with_cors(api):
    client, _ = api
    response = client.options("/api/uploads", headers={"Origin": "https://untrusted.invalid", "Access-Control-Request-Method": "POST"})
    assert response.status_code == 400
    assert "access-control-allow-origin" not in response.headers


def test_workbook_multisheet_reordered_headers_and_source_rows(api):
    book = Workbook()
    book.active.title = "说明"
    book.active.append(["说明页"])
    sheet = book.create_sheet("第一仓")
    sheet.append(["报表标题"])
    sheet.append([])
    sheet.append(list(reversed(HEADERS)))
    sheet.append(list(reversed(sample_row())))
    sheet.append([])
    sheet.append(list(reversed(sample_row("000456", "+ 2"))))
    second = book.create_sheet("第二仓")
    second.append(HEADERS)
    second.append(sample_row("000789", "0"))
    output = BytesIO()
    book.save(output)
    client, _ = api
    response = client.post("/api/uploads", headers=auth(client), files={"file": ("sheets.xlsx", output.getvalue())})
    assert response.status_code == 201
    assert response.json()["recordCount"] == 3
    assert response.json()["sheetCount"] == 2
    assert response.json()["skippedSheets"][0]["name"] == "说明"
    shared = client.get("/api/uploads/" + response.json()["id"]).json()
    assert shared["sheets"][0]["rowNumbers"] == [4, 6]
    assert shared["sheets"][0]["rows"][0][0] == "000123"
    assert shared["sheets"][1]["rows"][0][0] == "000789"


def test_50000_record_boundary_and_total_workbook_limit():
    book = Workbook()
    sheet = book.active
    sheet.append(HEADERS)
    for _ in range(MAX_RECORDS):
        sheet.append(sample_row())
    output = BytesIO()
    book.save(output)
    assert parse_excel(output.getvalue())["recordCount"] == MAX_RECORDS
    sheet.append(sample_row())
    output = BytesIO()
    book.save(output)
    with pytest.raises(ValueError, match="50,000"):
        parse_excel(output.getvalue())


def test_late_insert_failure_rolls_back_first_batch(api):
    client, repo = api
    calls = 0

    def fail_second_batch(connection, cursor, statement, parameters, context, executemany):
        nonlocal calls
        if statement.startswith("INSERT INTO pages_inventory_records"):
            calls += 1
            if calls == 2:
                raise RuntimeError("isolated second batch failure")

    event.listen(repo.engine, "before_cursor_execute", fail_second_batch)
    try:
        content = workbook_bytes([sample_row(str(index)) for index in range(501)])
        response = client.post("/api/uploads", headers=auth(client), files={"file": ("batch.xlsx", content)})
    finally:
        event.remove(repo.engine, "before_cursor_execute", fail_second_batch)
    assert response.status_code == 503
    assert calls == 2
    with repo.engine.connect() as connection:
        assert connection.scalar(select(func.count()).select_from(uploads)) == 0
        assert connection.scalar(select(func.count()).select_from(records)) == 0
    assert client.post("/api/uploads", headers=auth(client), files={"file": ("batch.xlsx", content)}).status_code == 201


def test_parallel_duplicate_imports_only_publish_once(tmp_path):
    engine = create_engine("sqlite:///" + (tmp_path / "concurrent.sqlite3").as_posix(), connect_args={"timeout": 30})
    repo = InventoryRepository(engine)
    repo.initialize()
    content = workbook_bytes([sample_row("000001"), sample_row("000002")])
    parsed = parse_excel(content)
    try:
        with ThreadPoolExecutor(max_workers=2) as workers:
            results = list(workers.map(lambda name: repo.save_upload(name, content, parsed), ["first.xlsx", "second.xlsx"]))
        assert results[0]["id"] == results[1]["id"]
        with engine.connect() as connection:
            assert connection.scalar(select(func.count()).select_from(uploads)) == 1
            assert connection.scalar(select(func.count()).select_from(records)) == 2
    finally:
        engine.dispose()


def test_unready_database_returns_503_instead_of_serving_empty_data():
    class UnavailableRepository:
        engine = type("DisposableEngine", (), {"dispose": lambda self: None})()

        def initialize(self):
            raise RuntimeError("isolated initialization failure")

    app = create_app(UnavailableRepository(), Settings(Path("unused.yaml"), PASSWORD, []))
    with TestClient(app) as client:
        for endpoint in ["/api/health", "/api/uploads", "/api/uploads/00000000-0000-4000-8000-000000000001"]:
            response = client.get(endpoint)
            assert response.status_code == 503
            assert "isolated initialization failure" not in response.text
